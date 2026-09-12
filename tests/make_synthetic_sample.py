"""Build a synthetic client-style sample WITHOUT any real client data.

Mirrors the structure the real minutes use (verified against a private sample):
  - title paragraph with soft line breaks
  - a 3-column attendees table (Sr.# | Name | Designation)
  - a 6-column action table (Sr.# | Agenda Items | AP | Department | Dead line | Remarks)
  - a header with an embedded logo image
  - a pre-existing 'Agenda Points:' section

This lets the golden suite run in CI / on a fresh clone with no client material.
"""
import io
import os

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _png_bytes():
    """Minimal valid 1x1 PNG (no external deps)."""
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAF"
        "AAH/q842iQAAAABJRU5ErkJggg=="
    )


def _add_logo(doc, stream):
    header = doc.sections[0].header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(stream, width=Inches(1.2))


def _table(doc, headers, rows, style="Table Grid"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = style
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(11)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            t.rows[ri].cells[ci].text = str(val)
    return t


def build(path):
    """Write a synthetic sample to `path` and return the path."""
    doc = Document()

    # --- header logo ---------------------------------------------------------
    logo = io.BytesIO(_png_bytes())
    _add_logo(doc, logo)

    # --- title block with soft line breaks (mimics the real sample) ----------
    p = doc.add_paragraph()
    r = p.add_run("MEETING MINUTES OF PROGRESS REVIEW")
    r.bold = True
    r.font.size = Pt(14)
    p.add_run().add_break()
    p.add_run("Date: 11-09-2026")
    p.add_run().add_break()
    p.add_run("Time: 10:00 AM")
    p.add_run().add_break()
    p.add_run("Minute Taker: P. Sample")

    doc.add_paragraph("ATTENDEES")
    _table(doc, ["Sr.#", "Name", "Designation"],
           [[1, "PERSON ONE", "Chief Executive"],
            [2, "PERSON TWO", "Director"],
            [3, "PERSON THREE", "Deputy Director"]])

    doc.add_paragraph("Agenda Points:")
    doc.add_paragraph("1. REVIEW OF PROGRESS")
    doc.add_paragraph("Discussion/update on the tasks assigned;")
    doc.add_paragraph("2. VENDOR PERFORMANCE")
    doc.add_paragraph("3. NEXT MEETING")

    doc.add_paragraph("ACTION POINTS")
    _table(doc, ["Sr.#", "Agenda Items", "AP", "Department", "Dead line", "Remarks"],
           [[1, "Item one", "Old action one", "Admin", "13 Sep", "Sample remark"],
            [2, "Item two", "Old action two", "Operations", "TBD", ""],
            [3, "Item three", "Old action three", "Finance", "20 Sep", ""]])

    # --- appended repeater sections (the shape the client added later) --------
    # A heading followed by a placeholder paragraph standing in for a LIST. These
    # are the "action items section" and "transcript section" that must become
    # {%p for %} loops rather than being purged. Kept in the synthetic sample so
    # the golden suite can prove the behaviour on a fresh clone.
    doc.add_paragraph("ACTION ITEMS")
    doc.add_paragraph("<Action Items Table goes here>")

    doc.add_paragraph("4. CLOSING")
    doc.add_paragraph("The meeting was adjourned.")

    doc.add_paragraph("COMPLETE DIARIZED TRANSCRIPT OF MEETING WITH TRANSLATION:")
    doc.add_paragraph("The complete transcript with translation text goes here")

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_synthetic.docx")
    build(out)
    print("wrote", out)
